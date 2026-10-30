"""영수증 JSON -> 간편장부 엑셀 (장부 / 월별 합계 / 확인필요)
사용: python make_ledger.py receipts.json 장부.xlsx
"""
import json, sys
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ACCOUNTS = ["상품·원재료 매입", "복리후생비", "여비교통비", "차량유지비", "기업업무추진비", "소모품비",
            "사무용품비", "도서인쇄비", "통신비", "광고선전비", "지급수수료", "임차료", "수도광열비", "교육훈련비", "기타"]
COLS = [("일자", "date", 12), ("거래처", "vendor", 18), ("거래내용", "item", 24), ("계정과목", "account", 16),
        ("금액", "amount", 12), ("부가세", "vat", 10), ("결제수단", "pay", 10), ("카드 끝4자리", "card_last4", 11),
        ("증빙", "proof", 12), ("확인필요", "check", 26)]
HEAD = PatternFill("solid", fgColor="1F2937")
WARN = PatternFill("solid", fgColor="FFF4CC")
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    s = "".join(ch for ch in str(v) if ch.isdigit() or ch == "-")
    return int(s) if s not in ("", "-") else None


def header(ws, names, widths):
    for i, (n, w) in enumerate(zip(names, widths), 1):
        c = ws.cell(row=1, column=i, value=n)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def main(src, dst):
    rows = json.load(open(src, encoding="utf-8"))
    if isinstance(rows, dict):
        rows = rows.get("receipts", [])
    for r in rows:
        r["amount"] = num(r.get("amount"))
        r["vat"] = num(r.get("vat"))
        if r.get("account") not in ACCOUNTS:
            r["check"] = ((r.get("check") or "") + " 계정과목 확인").strip()
            r["account"] = r.get("account") or "기타"
        cl = str(r.get("card_last4") or "")
        r["card_last4"] = cl[-4:] if cl.isdigit() else ""
    rows.sort(key=lambda r: (r.get("date") or "9999", r.get("vendor") or ""))

    wb = Workbook()
    ws = wb.active
    ws.title = "장부"
    header(ws, [c[0] for c in COLS], [c[2] for c in COLS])
    for i, r in enumerate(rows, 2):
        for j, (_, k, _) in enumerate(COLS, 1):
            c = ws.cell(row=i, column=j, value=r.get(k))
            c.border = BOX
            if k in ("amount", "vat"):
                c.number_format = "#,##0"
            if r.get("check"):
                c.fill = WARN
    n = len(rows) + 1
    t = n + 1
    ws.cell(row=t, column=4, value="합계").font = Font(bold=True)
    c = ws.cell(row=t, column=5, value=f"=SUM(E2:E{n})" if rows else 0)
    c.number_format = "#,##0"
    c.font = Font(bold=True)
    c = ws.cell(row=t, column=6, value=f"=SUM(F2:F{n})" if rows else 0)
    c.number_format = "#,##0"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{n}"

    # 월별 합계 (장부 시트 수식 참조)
    ms = wb.create_sheet("월별 합계")
    months = sorted({(r.get("date") or "")[:7] for r in rows if r.get("date")})
    used = [a for a in ACCOUNTS if any(r.get("account") == a for r in rows)]
    header(ms, ["계정과목"] + months + ["합계"], [18] + [12] * len(months) + [13])
    for i, a in enumerate(used, 2):
        ms.cell(row=i, column=1, value=a).border = BOX
        for j, m in enumerate(months, 2):
            f = f'=SUMIFS(장부!$E$2:$E${n},장부!$D$2:$D${n},$A{i},장부!$A$2:$A${n},"{m}*")'
            c = ms.cell(row=i, column=j, value=f)
            c.number_format = "#,##0"
            c.border = BOX
        c = ms.cell(row=i, column=len(months) + 2, value=f"=SUM(B{i}:{get_column_letter(len(months) + 1)}{i})")
        c.number_format = "#,##0"
        c.border = BOX
        c.font = Font(bold=True)
    tr = len(used) + 2
    ms.cell(row=tr, column=1, value="월 합계").font = Font(bold=True)
    for j in range(2, len(months) + 3):
        col = get_column_letter(j)
        c = ms.cell(row=tr, column=j, value=f"=SUM({col}2:{col}{tr - 1})")
        c.number_format = "#,##0"
        c.font = Font(bold=True)

    # 장부 날짜가 문자열이어도 SUMIFS 와일드카드가 맞도록 텍스트로 둔다
    for i in range(2, n + 1):
        ws.cell(row=i, column=1).number_format = "@"

    # 확인필요
    cs = wb.create_sheet("확인필요")
    header(cs, ["일자", "거래처", "금액", "확인필요"], [12, 18, 12, 40])
    k = 2
    for r in rows:
        if r.get("check"):
            for j, key in enumerate(["date", "vendor", "amount", "check"], 1):
                c = cs.cell(row=k, column=j, value=r.get(key))
                c.border = BOX
                if key == "amount":
                    c.number_format = "#,##0"
            k += 1

    wb.save(dst)
    total = sum(r["amount"] or 0 for r in rows)
    print(f"영수증 {len(rows)}장 -> 장부 {len(rows)}줄, 합계 {total:,}원, 확인 필요 {k - 2}건 -> {dst}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "장부.xlsx")
