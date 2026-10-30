"""장부 엑셀 vs 카드 이용내역 대조 -> 빠진 경비 후보 찾기
사용: python match_card.py 장부.xlsx 카드내역.(xlsx|csv) 대조결과.xlsx [--days 3]
카드 내역의 열 이름은 자동으로 찾는다(이용일/거래일/승인일, 가맹점명, 이용금액 등).
"""
import csv, sys, re
from datetime import datetime, date, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DATE_KEYS = ["이용일자", "이용일", "거래일자", "거래일", "승인일자", "승인일", "사용일", "일자", "날짜", "date"]
SHOP_KEYS = ["가맹점명", "이용가맹점", "가맹점", "이용처", "사용처", "상호", "거래처", "merchant"]
AMT_KEYS = ["이용금액", "승인금액", "결제금액", "사용금액", "금액", "원화금액", "amount"]
CANCEL_KEYS = ["취소", "승인취소", "상태", "구분"]


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    m = re.search(r"(\d{2,4})[.\-/년 ]+(\d{1,2})[.\-/월 ]+(\d{1,2})", s)
    if not m:
        m2 = re.fullmatch(r"(\d{4})(\d{2})(\d{2})", s[:8])
        if not m2:
            return None
        m = m2
    y, mo, d = (int(x) for x in m.groups())
    if y < 100:
        y += 2000
    try:
        return date(y, mo, d)
    except ValueError:
        return None


def to_amt(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    s = "".join(ch for ch in str(v) if ch.isdigit() or ch == "-")
    return int(s) if s not in ("", "-") else None


def read_table(path):
    if path.lower().endswith(".csv"):
        for enc in ("utf-8-sig", "cp949", "euc-kr"):
            try:
                return list(csv.reader(open(path, encoding=enc)))
            except UnicodeDecodeError:
                continue
        raise SystemExit("CSV 인코딩을 읽지 못했어요")
    ws = load_workbook(path, data_only=True).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def find_col(head, keys):
    h = [str(x or "").replace(" ", "") for x in head]
    for k in keys:
        for i, x in enumerate(h):
            if k.lower() in x.lower():
                return i
    return None


def load_card(path):
    rows = read_table(path)
    for hi, head in enumerate(rows[:15]):
        di, si, ai = find_col(head, DATE_KEYS), find_col(head, SHOP_KEYS), find_col(head, AMT_KEYS)
        if di is not None and ai is not None:
            ci = find_col(head, CANCEL_KEYS)
            out = []
            for r in rows[hi + 1:]:
                if not r or di >= len(r):
                    continue
                d, a = to_date(r[di]), to_amt(r[ai] if ai < len(r) else None)
                if not d or not a:
                    continue
                if ci is not None and ci < len(r) and "취소" in str(r[ci] or ""):
                    continue
                out.append({"date": d, "shop": str(r[si] if si is not None and si < len(r) else "").strip(), "amount": a})
            return out
    raise SystemExit("카드 내역에서 날짜·금액 열을 찾지 못했어요. 어느 열이 날짜·가맹점·금액인지 알려주세요.")


def load_ledger(path):
    ws = load_workbook(path, data_only=True)["장부"]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0] or r[3] == "합계":
            continue
        d, a = to_date(r[0]), to_amt(r[4])
        if d and a:
            out.append({"date": d, "shop": r[1] or "", "item": r[2] or "", "account": r[3] or "", "amount": a})
    return out


def sheet(wb, title, head, rows, fill=None):
    ws = wb.create_sheet(title)
    for i, h in enumerate(head, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = 16 if i != 2 else 22
    for k, r in enumerate(rows, 2):
        for i, v in enumerate(r, 1):
            c = ws.cell(row=k, column=i, value=v)
            if isinstance(v, int):
                c.number_format = "#,##0"
            if fill:
                c.fill = PatternFill("solid", fgColor=fill)
    ws.freeze_panes = "A2"


def main(ledger_path, card_path, out_path, days=3):
    led, card = load_ledger(ledger_path), load_card(card_path)
    used = set()
    matched, only_ledger = [], []
    for L in led:
        best = None
        for j, C in enumerate(card):
            if j in used or C["amount"] != L["amount"]:
                continue
            gap = abs((C["date"] - L["date"]).days)
            if gap <= days and (best is None or gap < best[1]):
                best = (j, gap)
        if best:
            used.add(best[0])
            C = card[best[0]]
            matched.append([L["date"].isoformat(), L["shop"], L["amount"], C["date"].isoformat(), C["shop"]])
        else:
            only_ledger.append([L["date"].isoformat(), L["shop"], L["amount"], L["account"]])
    missing = [[C["date"].isoformat(), C["shop"], C["amount"], "사업용이면 장부에 추가"] for j, C in enumerate(card) if j not in used]

    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "장부에 없음", ["카드 이용일", "가맹점", "금액", "할 일"], missing, "FFE4E6")
    sheet(wb, "카드에 없음", ["장부 일자", "거래처", "금액", "계정과목"], only_ledger, "FFF4CC")
    sheet(wb, "매칭됨", ["장부 일자", "거래처", "금액", "카드 이용일", "가맹점"], matched)
    wb.save(out_path)
    tot = sum(r[2] for r in missing)
    print(f"카드 {len(card)}건 · 장부 {len(led)}건 -> 매칭 {len(matched)}, 빠진 경비 후보 {len(missing)}건 (합계 {tot:,}원), 카드에 없음 {len(only_ledger)}건 -> {out_path}")
    for r in sorted(missing, key=lambda r: -r[2])[:5]:
        print(f"  - {r[0]} {r[1]} {r[2]:,}원")


if __name__ == "__main__":
    a = [x for x in sys.argv[1:] if not x.startswith("--")]
    d = 3
    if "--days" in sys.argv:
        d = int(sys.argv[sys.argv.index("--days") + 1])
    main(a[0], a[1], a[2] if len(a) > 2 else "대조결과.xlsx", d)
